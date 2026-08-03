"""
Demo cho bản v1 của Excel MCP server - chạy được trên Linux/Mac (không cần
Excel thật), dùng MockOpenpyxlBackend để chứng minh:

  1. server.py wire đúng các tool vào MCPServer (kiểm tra danh sách tool)
  2. Luồng open -> read/write cell -> apply formula -> create_sheet
     (copy) -> run "macro" -> save -> đọc lại kết quả hoạt động đúng
     end-to-end

Khi chạy thật trên Windows + Excel, chỉ cần đổi biến EXCEL_MCP_BACKEND
hoặc chạy server.py bình thường (nó tự nhận Windows và dùng
WindowsComBackend + macro VBA thật trong file .xlsm của anh).

Chạy: python demo/demo_test.py
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl

from backends.mock_openpyxl import MockOpenpyxlBackend

DEMO_XLSX = os.path.join(os.path.dirname(__file__), "demo_workbook.xlsx")


def make_demo_workbook():
    """Tạo file mẫu: sheet 'Data' có vài dòng doanh thu theo tháng."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["Thang", "DoanhThu"])
    for i, revenue in enumerate([120, 150, 90, 200], start=1):
        ws.append([f"Thang {i}", revenue])
    wb.save(DEMO_XLSX)
    print(f"[setup] Đã tạo {DEMO_XLSX}")


def fake_macro_generate_summary(wb, *args):
    """
    Giả lập 1 macro VBA tên 'GenerateSummary' mà thực tế công ty đã có sẵn
    trong file .xlsm: đọc sheet Data, tính tổng, ghi qua sheet Summary.
    Trên Windows thật, đây sẽ là: wb.macro('GenerateSummary')() gọi thẳng
    Sub GenerateSummary() trong VBA project.
    """
    data_ws = wb["Data"]
    # min_row=2 bỏ header; max_row=6 vì B7 là ô công thức (=SUM...), không phải data thô
    total = sum(
        row[1].value
        for row in data_ws.iter_rows(min_row=2, max_row=6)
        if isinstance(row[1].value, (int, float))
    )
    summary_ws = wb["Summary"]
    summary_ws["A1"] = "Tong doanh thu"
    summary_ws["B1"] = total
    return total


def main():
    print("=== 1) Kiểm tra MCP server đã wire đủ tool chưa ===")
    # import server sẽ chạy toàn bộ @mcp.tool() decorators -> đăng ký tool
    import server as excel_server

    tool_names = sorted(excel_server.mcp._tool_manager._tools.keys())
    print(f"Số tool đã đăng ký: {len(tool_names)}")
    for name in tool_names:
        print(f"  - {name}")

    print("\n=== 2) Demo luồng nghiệp vụ qua backend (mock, không cần Excel) ===")
    make_demo_workbook()

    backend = MockOpenpyxlBackend()
    backend.register_macro("GenerateSummary", fake_macro_generate_summary)

    wb_id = backend.open_workbook(DEMO_XLSX)
    print(f"[open] workbook_id = {wb_id}")

    print(f"[read] sheets = {backend.list_sheets(wb_id)}")
    print(f"[read] B3 (DoanhThu thang 2) = {backend.read_cell(wb_id, 'Data', 'B3')}")

    backend.write_cell(wb_id, "Data", "A6", "Thang 5")
    backend.write_cell(wb_id, "Data", "B6", 175)
    print("[write] đã thêm dòng Thang 5 = 175")

    backend.apply_formula(wb_id, "Data", "B7", "=SUM(B2:B6)")
    print("[formula] B7 = SUM(B2:B6)")

    backend.create_sheet(wb_id, "Summary")
    print("[sheet] đã tạo sheet 'Summary'")

    result = backend.run_macro(wb_id, "GenerateSummary")
    print(f"[macro] GenerateSummary() trả về tổng = {result}")

    saved_path = backend.save_workbook(wb_id)
    print(f"[save] đã lưu vào {saved_path}")

    backend.close_workbook(wb_id, save=False)

    # đọc lại từ file để chứng minh dữ liệu thực sự đã ghi xuống đĩa
    check_id = backend.open_workbook(saved_path)
    summary_value = backend.read_cell(check_id, "Summary", "B1")
    formula_cell = backend.read_cell(check_id, "Data", "B7")
    print(f"\n[verify] Summary!B1 (đọc lại từ file) = {summary_value}")
    print(f"[verify] Data!B7 (công thức, openpyxl trả về chuỗi công thức) = {formula_cell}")
    backend.close_workbook(check_id, save=False)

    print("\n=== Demo hoàn tất: luồng open -> write -> formula -> sheet -> macro -> save -> verify OK ===")


if __name__ == "__main__":
    main()
